"""V-wave: the postal/region case — T005S membership, postal formats,
placeholders, reference-less degradation."""
from openpyxl import Workbook

from mdmdoc.bulk import region


def _t005s(tmp_path):
    p = tmp_path / "t005s.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Country/Region Key", "Region", "Provincial Tax Code",
               "State of manufacture", "Description"])
    for r in (["US", "TX", "", "", "Texas"], ["US", "NY", "", "", "New York"],
              ["DE", "BY", "", "", "Bayern"], ["TH", "98", "", "", "Bangkok"]):
        ws.append(r)
    wb.save(p)
    return p


def _row(**kw):
    base = {"partner": "1", "country": "US", "region": "TX",
            "postal_code": "75201", "city": "Dallas"}
    base.update(kw)
    return base


def test_valid_row_with_reference(tmp_path):
    out = region.check_rows([_row()], refs=[_t005s(tmp_path)])
    assert out[0].bucket == "VALID"
    assert any("Texas" in r for r in out[0].reasons)


def test_unknown_region_invalid(tmp_path):
    out = region.check_rows([_row(region="ZZ")], refs=[_t005s(tmp_path)])
    assert out[0].bucket == "INVALID" and "BULK-R02" in out[0].rule_ids


def test_empty_region_required_vs_optional(tmp_path):
    ref = [_t005s(tmp_path)]
    out = region.check_rows([_row(region="")], refs=ref)
    assert out[0].bucket == "INVALID" and "BULK-R03" in out[0].rule_ids
    out = region.check_rows([_row(country="DE", region="", postal_code="80331")],
                            refs=ref)
    assert out[0].bucket == "SUSPICIOUS"


def test_postal_formats():
    out = region.check_rows([_row(country="DE", region="", postal_code="1234")])
    assert "BULK-R04" in out[0].rule_ids
    out = region.check_rows([_row(country="AE", region="", postal_code="12345")])
    assert "BULK-R05" in out[0].rule_ids
    out = region.check_rows([_row(country="CA", region="", postal_code="K1A 0B1")])
    assert "BULK-R04" not in out[0].rule_ids


def test_placeholders_and_empty():
    out = region.check_rows([_row(region="Foreign", postal_code="99")])
    assert out[0].bucket == "SUSPICIOUS" and "BULK-R06" in out[0].rule_ids
    out = region.check_rows([{"partner": "1", "country": "", "region": "",
                              "postal_code": "", "city": ""}])
    assert out[0].bucket == "SKIPPED" and "BULK-R07" in out[0].rule_ids


def test_no_reference_notes_degradation():
    notes: list = []
    out = region.check_rows([_row()], refs=[], notes=notes)
    assert out[0].bucket == "VALID"                 # postal format only
    assert any("no T005S" in n for n in notes)
